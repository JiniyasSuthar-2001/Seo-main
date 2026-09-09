import io
import csv
from typing import Dict, Any, List, Optional

from app.config.utils import sanitize_csv_cell
from app.services.crawl_data.crawl_dataset_service import CrawlDatasetService

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    openpyxl = None
    Font = PatternFill = Alignment = Border = Side = get_column_letter = None
    HAS_OPENPYXL = False


class CrawlTabExportService:
    """
    Export generator for Screaming-Frog-style crawl data.
    Provides complete CSV string generation and multi-worksheet XLSX workbooks.
    """

    ALL_TABS = [
        "internal",
        "response-codes",
        "titles",
        "meta-descriptions",
        "h1",
        "h2",
        "images",
        "canonicals",
        "directives",
        "hreflang",
        "structured-data",
        "redirects",
        "internal-links",
        "external-links",
        "broken-links",
        "issues"
    ]

    SHEET_NAMES = {
        "internal": "Internal",
        "response-codes": "Response Codes",
        "titles": "Page Titles",
        "meta-descriptions": "Meta Descriptions",
        "h1": "H1",
        "h2": "H2",
        "images": "Images",
        "canonicals": "Canonicals",
        "directives": "Directives",
        "hreflang": "Hreflang",
        "structured-data": "Structured Data",
        "redirects": "Redirects",
        "internal-links": "Internal Links",
        "external-links": "External Links",
        "broken-links": "Broken Links",
        "issues": "Issues"
    }

    @classmethod
    def export_tab_csv(
        cls,
        project_id: str,
        domain: Optional[str] = None,
        tab_name: str = "internal",
        crawl_id: Optional[str] = None
    ) -> str:
        """
        Generates full CSV content for a specific tab across the entire crawl dataset.
        All cell values are sanitized against CSV injection attacks.
        """
        artifacts = CrawlDatasetService.load_crawl_artifacts(project_id, domain, crawl_id)
        if not artifacts:
            return "No crawl data available\n"

        clean_tab = tab_name.strip().lower().replace("_", "-")
        meta = CrawlDatasetService.TAB_METADATA.get(clean_tab) or {
            "columns": []
        }
        columns = meta.get("columns", [])
        rows = CrawlDatasetService.get_tab_rows(artifacts, clean_tab)

        output = io.StringIO()
        writer = csv.writer(output)

        # Write Header
        header_labels = [c["label"] for c in columns] if columns else (list(rows[0].keys()) if rows else ["URL"])
        writer.writerow([sanitize_csv_cell(h) for h in header_labels])

        # Write Data Rows
        if columns:
            col_keys = [c["key"] for c in columns]
            for r in rows:
                row_vals = [sanitize_csv_cell(r.get(k, "")) for k in col_keys]
                writer.writerow(row_vals)
        else:
            for r in rows:
                row_vals = [sanitize_csv_cell(v) for v in r.values()]
                writer.writerow(row_vals)

        return output.getvalue()

    @classmethod
    def export_all_tabs_csv_dict(
        cls,
        project_id: str,
        domain: Optional[str] = None,
        crawl_id: Optional[str] = None
    ) -> Dict[str, str]:
        """
        Returns a mapping of { filename: csv_content } for all 16 tabs to embed into ZIP archives.
        """
        artifacts = CrawlDatasetService.load_crawl_artifacts(project_id, domain, crawl_id)
        if not artifacts:
            return {}

        result = {}
        for tab in cls.ALL_TABS:
            csv_content = cls.export_tab_csv(project_id, domain, tab, crawl_id)
            result[f"{tab}.csv"] = csv_content

        return result

    @classmethod
    def export_crawl_data_xlsx(
        cls,
        project_id: str,
        domain: Optional[str] = None,
        crawl_id: Optional[str] = None,
        project_name: Optional[str] = None
    ) -> bytes:
        """
        Generates a comprehensive, professionally styled Excel workbook (SEO_Crawl_Data.xlsx)
        containing all 16 Screaming-Frog-style crawl data worksheets.
        """
        if not HAS_OPENPYXL:
            raise RuntimeError("openpyxl library is required for XLSX export.")

        artifacts = CrawlDatasetService.load_crawl_artifacts(project_id, domain, crawl_id)
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style definitions
        FONT_FAMILY = "Arial"
        header_fill = PatternFill(start_color="0D2F5E", end_color="0D2F5E", fill_type="solid")
        header_font = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
        data_font = Font(name=FONT_FAMILY, size=9)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )
        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")

        if not artifacts:
            # Create a single placeholder sheet if no crawl data exists
            ws = wb.create_sheet(title="No Data")
            ws["A1"] = "No completed crawl data available for this website."
            ws["A1"].font = Font(name=FONT_FAMILY, size=11, bold=True)
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output.getvalue()

        for tab_key in cls.ALL_TABS:
            sheet_title = cls.SHEET_NAMES.get(tab_key, tab_key[:31])
            # Excel limits sheet names to 31 chars
            sheet_title = sheet_title[:31]
            ws = wb.create_sheet(title=sheet_title)
            ws.views.sheetView[0].showGridLines = True

            meta = CrawlDatasetService.TAB_METADATA.get(tab_key) or {"columns": []}
            columns = meta.get("columns", [])
            rows = CrawlDatasetService.get_tab_rows(artifacts, tab_key)

            col_keys = [c["key"] for c in columns]
            col_labels = [c["label"] for c in columns]

            # 1. Write Header Row
            for col_idx, label in enumerate(col_labels, 1):
                cell = ws.cell(row=1, column=col_idx, value=label)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_left
                cell.border = thin_border
            ws.row_dimensions[1].height = 24

            # 2. Write Data Rows
            for row_idx, r in enumerate(rows, 2):
                ws.row_dimensions[row_idx].height = 18
                for col_idx, k in enumerate(col_keys, 1):
                    val = r.get(k, "")
                    if val is None:
                        val = ""
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.font = data_font
                    cell.border = thin_border
                    # Numeric / short text alignment
                    if isinstance(val, (int, float)) or str(val).lower() in ("yes", "no", "2xx", "3xx", "4xx", "5xx", "critical", "warning"):
                        cell.alignment = align_center
                    else:
                        cell.alignment = align_left

            # 3. Auto-fit column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 65)

            # Freeze header pane
            ws.freeze_panes = "A2"

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
